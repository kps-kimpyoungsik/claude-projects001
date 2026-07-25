"""XLSX 파서 어댑터(xlsx_adapter.py) 테스트.

openpyxl로 합성 워크북(시트 2개, 표 데이터)을 만들어 실제 parse_to_markdown() 호출
결과를 검증한다(placeholder 반환 금지, T98 AIP).
"""

import io

from openpyxl import Workbook

from backend.adapters.parsers.xlsx_adapter import XlsxParserAdapter


def _build_sample_xlsx() -> io.BytesIO:
    wb = Workbook()
    sheet1 = wb.active
    sheet1.title = "요구사항"
    sheet1.append(["번호", "항목"])
    sheet1.append(["1", "결제 시스템 구축"])
    sheet1.append(["2", "보안 요건 검토"])

    sheet2 = wb.create_sheet("빈시트")  # 빈 시트는 출력에서 생략되어야 함

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_can_handle_xlsx_only():
    adapter = XlsxParserAdapter()
    assert adapter.can_handle(".xlsx") is True
    assert adapter.can_handle(".XLSX") is True
    assert adapter.can_handle(".xls") is False
    assert adapter.can_handle(".docx") is False


def test_parse_to_markdown_extracts_sheet_as_table():
    adapter = XlsxParserAdapter()
    md = adapter.parse_to_markdown(_build_sample_xlsx(), metadata={"filename": "test.xlsx"})

    assert "## Sheet: 요구사항" in md
    assert "| 번호 | 항목 |" in md
    assert "| 1 | 결제 시스템 구축 |" in md
    assert "| 2 | 보안 요건 검토 |" in md
    # 빈 시트는 헤딩·표 없이 통째로 생략된다(빈 청크 방지)
    assert "빈시트" not in md


def test_parse_to_markdown_empty_workbook_returns_empty_string():
    wb = Workbook()
    wb.active.title = "빈시트"
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)

    adapter = XlsxParserAdapter()
    md = adapter.parse_to_markdown(buf, metadata={"filename": "empty.xlsx"})
    assert md == ""
