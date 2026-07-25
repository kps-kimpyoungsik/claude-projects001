"""PPTX 파서 어댑터(pptx_adapter.py) 테스트.

python-pptx로 합성 프레젠테이션(제목+2줄 불릿+테이블+speaker notes)을 만들어
실제 parse_to_markdown() 호출 결과를 검증한다(placeholder 반환 금지, T98 AIP).
"""

import io

from pptx import Presentation

from backend.adapters.parsers.pptx_adapter import PptxParserAdapter


def _build_sample_pptx() -> io.BytesIO:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Test Title"
    body = slide.placeholders[1]
    body.text_frame.text = "First bullet"
    p = body.text_frame.add_paragraph()
    p.text = "Second bullet"
    slide.notes_slide.notes_text_frame.text = "This is a note"

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)
    return buf


def test_can_handle_pptx_only():
    adapter = PptxParserAdapter()
    assert adapter.can_handle(".pptx") is True
    assert adapter.can_handle(".PPTX") is True
    assert adapter.can_handle(".ppt") is False
    assert adapter.can_handle(".docx") is False


def test_parse_to_markdown_extracts_title_bullets_and_notes():
    adapter = PptxParserAdapter()
    md = adapter.parse_to_markdown(_build_sample_pptx(), metadata={"filename": "test.pptx"})

    assert "## Slide 1" in md
    assert "Test Title" in md
    assert "First bullet" in md
    assert "Second bullet" in md
    assert "Speaker notes: This is a note" in md


def test_parse_to_markdown_extracts_table():
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    rows, cols = 2, 2
    table_shape = slide.shapes.add_table(rows, cols, 0, 0, 4000000, 1000000)
    table = table_shape.table
    table.cell(0, 0).text = "Header1"
    table.cell(0, 1).text = "Header2"
    table.cell(1, 0).text = "Val1"
    table.cell(1, 1).text = "Val2"

    buf = io.BytesIO()
    prs.save(buf)
    buf.seek(0)

    adapter = PptxParserAdapter()
    md = adapter.parse_to_markdown(buf, metadata={"filename": "table.pptx"})

    assert "| Header1 | Header2 |" in md
    assert "| Val1 | Val2 |" in md
    # [2026-07-22] 표는 별도 서브헤딩(###)으로 슬라이드 본문과 분리된다(청킹 방법론 재검토).
    assert "### Table 1 (Slide 1)" in md
