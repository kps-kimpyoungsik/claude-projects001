"""[2026-07-22 청킹 방법론 재검토 후속] XLSX -> Markdown 어댑터 (openpyxl 사용).

`router.py`의 FORMAT_STRATEGY[".xlsx"] = "spreadsheet_parse" 설계를 실제 구현한다.
`openpyxl`은 이 환경에 이미 설치되어 있음을 실측 확인(pip show openpyxl) 후 사용 —
신규 의존성 설치 없이 재사용한다(CRZ).

시트마다 "## Sheet: {이름}" 헤딩으로 구분한다 — pptx/pdf 어댑터의 "표는 프로즈와
분리한다" 원칙이 엑셀에서는 자연히 성립한다(시트 = 표 그 자체이므로 별도 서브헤딩
분리가 불필요, 시트 헤딩이 곧 표 경계). 시트 하나가 너무 크면(행 수 과다) 청킹
단계(HeadingBoundarySplitter)가 시트 전체를 한 청크로 묶어버려 컨텍스트 윈도우를
넘길 수 있다는 점을 실측 인지하되, 행 단위 재분할은 이번 범위에서 다루지 않는다
(문서 하나당 시트 수·행 수가 실사용 규모에서 문제될 정도인지 실측 근거가 없어
과설계하지 않음 — 필요해지면 후속 작업으로 분리).
"""

from typing import Any, BinaryIO

from openpyxl import load_workbook

from backend.application.ports.parser_port import ParserPort


class XlsxParserAdapter(ParserPort):
    def can_handle(self, file_extension: str) -> bool:
        return file_extension.lower() in (".xlsx",)
        # .xls(구버전 바이너리 포맷)는 openpyxl이 지원하지 않아 제외

    def parse_to_markdown(self, file_stream: BinaryIO, metadata: dict[str, Any]) -> str:
        workbook = load_workbook(file_stream, data_only=True, read_only=True)
        lines: list[str] = []
        for sheet in workbook.worksheets:
            table_lines = self._sheet_to_markdown(sheet)
            if not table_lines:
                continue
            lines.append(f"## Sheet: {sheet.title}")
            lines.extend(table_lines)
            lines.append("")
        return "\n".join(lines).strip() + "\n" if lines else ""

    def _sheet_to_markdown(self, sheet) -> list[str]:
        rows = [
            ["" if cell is None else str(cell) for cell in row]
            for row in sheet.iter_rows(values_only=True)
        ]
        rows = [row for row in rows if any(cell.strip() for cell in row)]
        if not rows:
            return []
        out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"]
        for row in rows[1:]:
            out.append("| " + " | ".join(row) + " |")
        return out
