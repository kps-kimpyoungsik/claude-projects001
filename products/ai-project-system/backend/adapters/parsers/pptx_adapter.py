"""[Phase 2 후속] PPTX -> Markdown 어댑터 (실제 동작 구현, python-pptx 사용).

`router.py`의 FORMAT_STRATEGY[".pptx"] = "slide_parse"(슬라이드 단위 분할 + speaker
notes 포함) 설계를 실제 구현한다. `python-pptx`는 이 환경에 이미 설치되어 있음을
실측 확인(pip show python-pptx) 후 사용 — docx_adapter.py처럼 stdlib만으로 재구현하지
않는다(이미 검증된 실치 라이브러리를 재발명하지 않음, CRZ).

슬라이드마다 "## Slide N" 헤딩으로 구분한다 — heading_splitter.py(SPCEngine)가
헤딩 경계로 청킹하므로, 슬라이드 = 청크 경계가 자연스럽게 성립한다(설계 의도 그대로).

[2026-07-22 청킹 방법론 재검토 반영] 외부 최신 사례(RAG 청킹 2026 모범사례 — 표는 프로즈와
섞인 채로 두지 않고 별도 청크로 분리해야 함) 조사 결과, 표를 슬라이드 텍스트와 같은 청크에
인라인으로 묻어두던 기존 구현은 표가 큰 슬라이드에서 산문+표가 뒤섞인 청크를 만드는 문제가
있었다(실측: HeadingBoundarySplitter는 `#`로 시작하는 줄에서만 경계를 나누므로, 이전 코드는
표를 어떤 헤딩으로도 구분하지 않아 슬라이드 청크 안에 통째로 포함됐다). 표마다
"### Table N (Slide M)" 서브헤딩을 붙여 슬라이드 본문과 별도 청크로 분리한다 — 헤딩 레벨을
3(`###`)으로 둬 heading_path에 "Slide M > Table N" 계층이 그대로 남는다(문맥 손실 없음).
"""

from typing import Any, BinaryIO

from pptx import Presentation

from backend.application.ports.parser_port import ParserPort


class PptxParserAdapter(ParserPort):
    def can_handle(self, file_extension: str) -> bool:
        return file_extension.lower() in (".pptx",)
        # .ppt(구버전 바이너리 포맷)는 python-pptx가 지원하지 않아 제외

    def parse_to_markdown(self, file_stream: BinaryIO, metadata: dict[str, Any]) -> str:
        prs = Presentation(file_stream)
        lines: list[str] = []
        for idx, slide in enumerate(prs.slides, start=1):
            lines.append(f"## Slide {idx}")
            body_lines, table_blocks = self._slide_text_lines(slide)
            if body_lines:
                lines.extend(body_lines)
            notes_text = self._speaker_notes(slide)
            if notes_text:
                lines.append("")
                lines.append("> Speaker notes: " + notes_text)
            lines.append("")
            # 표는 슬라이드 본문과 섞지 않고 서브헤딩(###)으로 별도 청크 경계를 만든다
            # (2026-07-22 청킹 방법론 재검토 — 표-프로즈 혼합 청크 방지, docstring 참조).
            for table_idx, table_lines in enumerate(table_blocks, start=1):
                lines.append(f"### Table {table_idx} (Slide {idx})")
                lines.extend(table_lines)
                lines.append("")
        return "\n".join(lines).strip() + "\n" if lines else ""

    def _slide_text_lines(self, slide) -> tuple[list[str], list[list[str]]]:
        out: list[str] = []
        tables: list[list[str]] = []
        for shape in slide.shapes:
            if shape.has_table:
                tables.append(self._table_to_markdown(shape.table))
                continue
            if not shape.has_text_frame:
                continue
            for para in shape.text_frame.paragraphs:
                text = "".join(run.text for run in para.runs).strip()
                if text:
                    out.append(text)
        return out, tables

    def _table_to_markdown(self, table) -> list[str]:
        rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
        if not rows:
            return []
        out = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join("---" for _ in rows[0]) + " |"]
        for row in rows[1:]:
            out.append("| " + " | ".join(row) + " |")
        return out

    def _speaker_notes(self, slide) -> str:
        if not slide.has_notes_slide:
            return ""
        frame = slide.notes_slide.notes_text_frame
        return frame.text.strip() if frame is not None else ""
