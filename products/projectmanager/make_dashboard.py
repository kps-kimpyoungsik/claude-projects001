# -*- coding: utf-8 -*-
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.formatting.rule import DataBarRule
from openpyxl.utils import get_column_letter

P = r'D:\projects\products\projectmanager\통합테스트_대응업무분장_담당자_V1.0.xlsx'
OUT = r'D:\projects\products\projectmanager\통합테스트_대응업무분장_담당자_V1.1_대시보드.xlsx'
SRC = "'PC_IA(화면목록)'"
# ponytail: 5~1000행 고정 범위 — 행 추가돼도 자동 반영
rng = lambda col: f"{SRC}!${col}$5:${col}$1000"
A, B, C, J, L = (rng(x) for x in 'ABCJL')
NOKEY = object()   # 인자 미지정과 "2Depth 공백" 구분

wb = openpyxl.load_workbook(P)
src = wb['PC_IA(화면목록)']

# 원본에서 그룹 키 수집(순서 보존)
d1, pairs = [], []
for r in range(5, src.max_row + 1):
    if not src.cell(r, 1).value:
        continue
    a, b = src.cell(r, 2).value, src.cell(r, 3).value
    if a not in d1:
        d1.append(a)
    if (a, b) not in pairs:
        pairs.append((a, b))

AREAS = ['수행', '현업']
STATES = ['완료', '처리중', '대기']

del wb['대시보드']
ws = wb.create_sheet('대시보드', wb.sheetnames.index('제개정이력') + 1)

HDR = PatternFill('solid', fgColor='1F4E79')
SUB = PatternFill('solid', fgColor='DDEBF7')
GRP = PatternFill('solid', fgColor='F2F2F2')
WHT = Font(color='FFFFFF', bold=True)
BD = Border(*[Side('thin', color='BFBFBF')] * 4)
CEN = Alignment('center', 'center')

def put(r, c, v, fill=None, font=None, fmt=None, al=CEN):
    cell = ws.cell(r, c, v)
    cell.border = BD; cell.alignment = al
    if fill: cell.fill = fill
    if font: cell.font = font
    if fmt: cell.number_format = fmt
    return cell

def cnt(state=None, area=None, d1v=None, d2v=NOKEY):
    p = []
    if state: p += [J, f'"{state}"']
    if area:  p += [L, f'"{area}"']
    if d1v is not None: p += [B, f'"{d1v}"']
    if d2v is not NOKEY: p += [C, '""' if d2v is None else f'"{d2v}"']
    if not p:
        return f'=COUNTA({A})'
    return '=COUNTIFS(' + ','.join(p) + ')'

ws['A1'] = '통합테스트 진행 대시보드 (자동 집계)'
ws['A1'].font = Font(size=16, bold=True, color='1F4E79')
ws['A2'] = "원본: 'PC_IA(화면목록)' J열(완료/처리중/대기) · L열(테스트 영역) 변경 시 자동 반영"
ws['A2'].font = Font(size=9, color='808080')

# ── 표1: 테스트 영역별 요약
r0 = 4
put(r0, 1, '① 테스트 영역별 요약', font=Font(bold=True, size=12, color='1F4E79'), al=Alignment('left', 'center'))
r0 += 1
for i, h in enumerate(['테스트 영역', '전체', '완료', '처리중', '대기', '진척률']):
    put(r0, 1 + i, h, HDR, WHT)
for k, area in enumerate(AREAS + ['합계']):
    r = r0 + 1 + k
    tot = 'sum'
    if area == '합계':
        put(r, 1, '합계', SUB, Font(bold=True))
        put(r, 2, f'=SUM(B{r0+1}:B{r-1})', SUB, Font(bold=True))
        for i, s in enumerate(STATES):
            cl = get_column_letter(3 + i)
            put(r, 3 + i, f'=SUM({cl}{r0+1}:{cl}{r-1})', SUB, Font(bold=True))
    else:
        put(r, 1, area, GRP)
        put(r, 2, cnt(area=area))
        for i, s in enumerate(STATES):
            put(r, 3 + i, cnt(state=s, area=area))
    put(r, 6, f'=IF(B{r}=0,"",C{r}/B{r})', SUB if area == '합계' else None,
        Font(bold=True) if area == '합계' else None, '0.0%')
sum_end = r0 + len(AREAS) + 1
ws.conditional_formatting.add(f'F{r0+1}:F{sum_end}',
    DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='63BE7B'))

# ── 표2: 1 Depth × 테스트 영역
r0 = sum_end + 2
put(r0, 1, '② 1 Depth × 테스트 영역', font=Font(bold=True, size=12, color='1F4E79'), al=Alignment('left', 'center'))
r0 += 1
hdrs = ['1 Depth', '전체', '완료', '처리중', '대기', '진척률'] + [f'{a}-{s}' for a in AREAS for s in STATES]
for i, h in enumerate(hdrs):
    put(r0, 1 + i, h, HDR, WHT)
for k, v in enumerate(d1):
    r = r0 + 1 + k
    put(r, 1, v, GRP, al=Alignment('left', 'center'))
    put(r, 2, cnt(d1v=v))
    for i, s in enumerate(STATES):
        put(r, 3 + i, cnt(state=s, d1v=v))
    put(r, 6, f'=IF(B{r}=0,"",C{r}/B{r})', fmt='0.0%')
    for i, (a, s) in enumerate([(a, s) for a in AREAS for s in STATES]):
        put(r, 7 + i, cnt(state=s, area=a, d1v=v))
d1_end = r0 + len(d1)
r = d1_end + 1
put(r, 1, '합계', SUB, Font(bold=True), al=Alignment('left', 'center'))
for c in range(2, 6 + len(AREAS) * len(STATES) + 1):
    if c == 6:
        put(r, 6, f'=IF(B{r}=0,"",C{r}/B{r})', SUB, Font(bold=True), '0.0%'); continue
    cl = get_column_letter(c)
    put(r, c, f'=SUM({cl}{r0+1}:{cl}{d1_end})', SUB, Font(bold=True))
ws.conditional_formatting.add(f'F{r0+1}:F{d1_end}',
    DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='63BE7B'))

# ── 표3: 1~2 Depth 상세
r0 = r + 2
put(r0, 1, '③ 1~2 Depth 상세 (2 Depth 없는 화면은 "(직접)")',
    font=Font(bold=True, size=12, color='1F4E79'), al=Alignment('left', 'center'))
r0 += 1
hdrs = ['1 Depth', '2 Depth', '전체', '완료', '처리중', '대기', '진척률'] + [f'{a}-{s}' for a in AREAS for s in STATES]
for i, h in enumerate(hdrs):
    put(r0, 1 + i, h, HDR, WHT)
for k, (a, b) in enumerate(pairs):
    r = r0 + 1 + k
    put(r, 1, a, GRP, al=Alignment('left', 'center'))
    put(r, 2, b if b else '(직접)', al=Alignment('left', 'center'))
    put(r, 3, cnt(d1v=a, d2v=b))
    for i, s in enumerate(STATES):
        put(r, 4 + i, cnt(state=s, d1v=a, d2v=b))
    put(r, 7, f'=IF(C{r}=0,"",D{r}/C{r})', fmt='0.0%')
    for i, (ar, s) in enumerate([(x, y) for x in AREAS for y in STATES]):
        put(r, 8 + i, cnt(state=s, area=ar, d1v=a, d2v=b))
p_end = r0 + len(pairs)
r = p_end + 1
put(r, 1, '합계', SUB, Font(bold=True), al=Alignment('left', 'center'))
put(r, 2, '', SUB)
for c in range(3, 7 + len(AREAS) * len(STATES) + 1):
    if c == 7:
        put(r, 7, f'=IF(C{r}=0,"",D{r}/C{r})', SUB, Font(bold=True), '0.0%'); continue
    cl = get_column_letter(c)
    put(r, c, f'=SUM({cl}{r0+1}:{cl}{p_end})', SUB, Font(bold=True))
ws.conditional_formatting.add(f'G{r0+1}:G{p_end}',
    DataBarRule(start_type='num', start_value=0, end_type='num', end_value=1, color='63BE7B'))

ws.column_dimensions['A'].width = 18
ws.column_dimensions['B'].width = 22
for c in range(3, 15):
    ws.column_dimensions[get_column_letter(c)].width = 9
ws.freeze_panes = 'A5'
wb.save(OUT)
print('OK  1Depth=%d  pairs=%d  lastrow=%d' % (len(d1), len(pairs), p_end + 1))
